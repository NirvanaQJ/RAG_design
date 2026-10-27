from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
from typing import Optional
from pydantic import BaseModel

from database import get_db, Document
from auth import require_admin, User
from document_processor import extract_text, split_text_into_chunks, save_upload_file
from vector_store import vector_store
from rag_service import rag_service

router = APIRouter(prefix="/api/admin", tags=["管理员"])


class DocumentUpdateRequest(BaseModel):
    title: Optional[str] = None
    original_url: Optional[str] = None
    description: Optional[str] = None


class AnalyzeResponse(BaseModel):
    title: str
    description: str


@router.post("/documents/analyze", response_model=AnalyzeResponse)
async def analyze_document(
    file: UploadFile = File(...),
    current_user: User = Depends(require_admin),
):
    allowed_extensions = {"docx", "doc", "pdf", "txt", "png", "jpg", "jpeg", "bmp", "tiff", "gif"}
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in allowed_extensions:
        raise HTTPException(400, f"不支持的文件格式: {ext}")

    content = await file.read()
    file_path, file_ext = save_upload_file(content, file.filename)

    image_exts = {"png", "jpg", "jpeg", "bmp", "tiff", "gif"}

    if file_ext in image_exts:
        result = rag_service.analyze_image(file_path)
    else:
        try:
            extracted_text = extract_text(file_path, file_ext)
        except Exception as e:
            raise HTTPException(500, f"文档解析失败: {str(e)}")

        if not extracted_text.strip():
            raise HTTPException(400, "文档内容为空或无法提取文本")

        try:
            result = rag_service.analyze_document(extracted_text)
        except Exception:
            lines = extracted_text.strip().split("\n")
            first_line = lines[0].strip() if lines else "未命名通知"
            if len(first_line) > 50:
                first_line = first_line[:50] + "..."
            preview = extracted_text[:100].replace("\n", " ").strip()
            result = {"title": first_line, "description": preview}

    return AnalyzeResponse(title=result["title"], description=result["description"])


@router.post("/documents/upload")
async def upload_document(
    file: UploadFile = File(...),
    title: str = Form(...),
    original_url: str = Form(""),
    description: str = Form(""),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    allowed_extensions = {"docx", "doc", "pdf", "txt", "png", "jpg", "jpeg", "bmp", "tiff", "gif"}
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in allowed_extensions:
        raise HTTPException(400, f"不支持的文件格式: {ext}，支持: {', '.join(allowed_extensions)}")

    content = await file.read()
    file_path, file_ext = save_upload_file(content, file.filename)

    try:
        extracted_text = extract_text(file_path, file_ext)
    except Exception as e:
        raise HTTPException(500, f"文档解析失败: {str(e)}")

    if not extracted_text.strip():
        raise HTTPException(400, "文档内容为空或无法提取文本")

    chunks = split_text_into_chunks(extracted_text)

    doc = Document(
        title=title,
        file_type=file_ext,
        original_url=original_url,
        description=description,
        content=extracted_text,
        chunk_count=len(chunks),
        uploaded_by=current_user.id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    vector_store.add_chunks(doc.id, chunks)

    return {
        "id": doc.id,
        "title": doc.title,
        "file_type": doc.file_type,
        "chunk_count": doc.chunk_count,
        "message": "文档上传并处理成功",
    }


@router.get("/documents")
def list_documents(
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    docs = db.query(Document).order_by(Document.created_at.desc()).all()
    return [
        {
            "id": d.id,
            "title": d.title,
            "file_type": d.file_type,
            "original_url": d.original_url,
            "description": d.description,
            "chunk_count": d.chunk_count,
            "created_at": d.created_at.isoformat(),
        }
        for d in docs
    ]


@router.put("/documents/{doc_id}")
def update_document(
    doc_id: int,
    req: DocumentUpdateRequest,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404, "文档不存在")

    if req.title is not None:
        doc.title = req.title
    if req.original_url is not None:
        doc.original_url = req.original_url
    if req.description is not None:
        doc.description = req.description

    db.commit()
    return {"message": "文档信息更新成功"}


@router.delete("/documents/{doc_id}")
def delete_document(
    doc_id: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(404, "文档不存在")

    vector_store.delete_document_chunks(doc_id)
    db.delete(doc)
    db.commit()
    return {"message": "文档已删除"}


@router.get("/stats")
def get_stats(
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    doc_count = db.query(Document).count()
    vs_stats = vector_store.get_collection_stats()
    return {
        "total_documents": doc_count,
        "total_chunks": vs_stats["total_chunks"],
    }


@router.post("/screenshot")
async def capture_screenshot(
    current_user: User = Depends(require_admin),
):
    try:
        from PIL import ImageGrab
        import base64
        from io import BytesIO

        screenshot = ImageGrab.grab(all_screens=True)
        buffer = BytesIO()
        screenshot.save(buffer, format="PNG")
        buffer.seek(0)

        img_base64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        data_url = f"data:image/png;base64,{img_base64}"

        return {"data_url": data_url, "width": screenshot.width, "height": screenshot.height}
    except Exception as e:
        raise HTTPException(500, f"截图失败: {str(e)}")


class ScrapeRequest(BaseModel):
    url: str


@router.post("/scrape")
async def scrape_url(
    req: ScrapeRequest,
    current_user: User = Depends(require_admin),
):
    import httpx
    from bs4 import BeautifulSoup
    from docx import Document as DocxDocument
    import base64
    from io import BytesIO
    import re

    url = req.url.strip()
    if not url.startswith(("http://", "https://")):
        raise HTTPException(400, "请输入有效的网页链接（以 http:// 或 https:// 开头）")

    try:
        async with httpx.AsyncClient(timeout=30.0, follow_redirects=True) as client:
            response = await client.get(url, headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            })
            response.raise_for_status()
            html = response.text
    except httpx.HTTPStatusError as e:
        raise HTTPException(400, f"网页请求失败：HTTP {e.response.status_code}")
    except httpx.RequestError as e:
        raise HTTPException(400, f"无法访问该链接：{str(e)}")
    except Exception as e:
        raise HTTPException(500, f"爬取过程出错：{str(e)}")

    try:
        soup = BeautifulSoup(html, "lxml")
    except Exception:
        soup = BeautifulSoup(html, "html.parser")

    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "iframe"]):
        tag.decompose()

    body = soup.find("body")
    if body:
        text = body.get_text(separator="\n", strip=True)
    else:
        text = soup.get_text(separator="\n", strip=True)

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    cleaned_lines = []
    for line in lines:
        if len(line) < 3:
            continue
        if re.match(r"^[\s\d\W]+$", line):
            continue
        cleaned_lines.append(line)

    content = "\n\n".join(cleaned_lines)

    if not content.strip():
        raise HTTPException(400, "未能从该网页提取到有效文本内容")

    result = rag_service.analyze_document(content)
    title = result["title"]
    description = f"来源：{url}\n{result['description']}"

    safe_title = re.sub(r'[\\/*?:"<>|]', "_", title)
    if len(safe_title) > 80:
        safe_title = safe_title[:80]

    doc = DocxDocument()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"原文链接：{url}")
    doc.add_paragraph("")
    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            doc.add_paragraph(paragraph)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    docx_base64 = base64.b64encode(buffer.getvalue()).decode("utf-8")

    return {
        "title": title,
        "description": description,
        "file_name": f"{safe_title}.docx",
        "file_base64": docx_base64,
        "content_length": len(content),
    }
